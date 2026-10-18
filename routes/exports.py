"""Excel (.xlsx) exports."""
from datetime import date
from io import BytesIO

from flask import Blueprint, abort, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from models.models import Customer, Expense, Product, Sale, Supplier
from routes.auth import owner_required
from services import report_service as rs
from utils import parse_date

bp = Blueprint("exports", __name__)


def xlsx_response(name, headers, rows, title="Export"):
    wb = Workbook()
    ws = wb.active
    ws.title = title[:30]
    ws.append(headers)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1E293B")
    for r in rows:
        ws.append(list(r))
    for i, h in enumerate(headers, 1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:200] if len(r) >= i]) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(width, 50)
    ws.freeze_panes = "A2"
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"{name}-{date.today():%Y%m%d}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@bp.route("/export/<kind>.xlsx")
@owner_required
def export(kind):
    if kind == "products":
        ps = Product.query.order_by(Product.name).all()
        return xlsx_response(kind, ["ID", "Name", "SKU", "Barcode", "Category", "Brand", "Unit", "Purchase price", "Selling price", "Stock", "Min stock", "Stock value", "Status"],
                             [[p.id, p.name, p.sku, p.barcode, p.category.name if p.category else "", p.brand, p.unit, p.purchase_price, p.selling_price,
                               p.stock_qty, p.min_stock, p.stock_value, p.status] for p in ps], "Products")
    if kind == "customers":
        cs = Customer.query.order_by(Customer.name).all()
        return xlsx_response(kind, ["ID", "Name", "Phone", "Email", "Address", "Total credit", "Payments", "Outstanding"],
                             [[c.id, c.name, c.phone, c.email, c.address, c.total_credit, c.total_payments, c.balance] for c in cs], "Customers")
    if kind == "suppliers":
        ss = Supplier.query.order_by(Supplier.name).all()
        return xlsx_response(kind, ["ID", "Name", "Company", "Phone", "Email", "Address", "Purchased", "Paid", "Payable"],
                             [[s.id, s.name, s.company, s.phone, s.email, s.address, s.total_purchased, s.total_paid, s.payable] for s in ss], "Suppliers")
    if kind == "sales":
        try:
            start, end = parse_date(request.args.get("start")), parse_date(request.args.get("end"))
        except ValueError:
            start = end = None
        q = Sale.query.order_by(Sale.id)
        if start:
            q = q.filter(Sale.date >= start)
        if end:
            q = q.filter(Sale.date <= end)
        return xlsx_response(kind, ["Invoice", "Date", "Customer", "Method", "Subtotal", "Discount", "Returns", "Total", "Paid", "Remaining", "Profit", "Status"],
                             [[s.invoice_no, s.date, s.customer.name if s.customer else "Walk-in", s.payment_method, s.subtotal, s.discount,
                               s.returned_total, s.total, s.paid, s.remaining, s.profit, s.status] for s in q.all()], "Sales")
    if kind == "expenses":
        es = Expense.query.order_by(Expense.date).all()
        return xlsx_response(kind, ["Date", "Title", "Category", "Amount", "Description"],
                             [[e.date, e.title, e.category, e.amount, e.description] for e in es], "Expenses")
    abort(404)


@bp.route("/reports/xlsx")
@owner_required
def report_xlsx():
    kind = request.args.get("type", "sales")
    if kind not in rs.REPORT_TYPES:
        abort(404)
    try:
        start, end = rs.resolve_range(request.args.get("period", "monthly"), request.args.get("start"), request.args.get("end"))
    except ValueError:
        abort(400)
    rep = rs.get_report(kind, start, end)
    rows = [list(r) for r in rep["rows"]] + [[], ["Summary"]] + [[l, v] for l, v in rep["summary"]]
    return xlsx_response(f"{kind}-report", rep["headers"], rows, rep["title"])
